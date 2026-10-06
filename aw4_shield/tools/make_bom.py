import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

parts=json.load(open(sys.argv[1]))["parts"]
OUT=sys.argv[2]
# value -> (group, name, spec, example part, label on silkscreen, polarity note, assembly step)
CAT={
 "330":   ("Rezistory","Rezistor 330 Ω","0,25 W, 1 %, metalizovaný, DIN 0207 (montáž nastojato, rozteč 2,54)","libovolný výrobce","330","",1),
 "220":   ("Rezistory","Rezistor 220 Ω","0,25 W, 1 %, DIN 0207","libovolný výrobce","220","",1),
 "1k":    ("Rezistory","Rezistor 1 kΩ","0,25 W, 1 %, DIN 0207","libovolný výrobce","1k","",1),
 "2k2":   ("Rezistory","Rezistor 2,2 kΩ","0,25 W, 1 %, DIN 0207","libovolný výrobce","2k2","",1),
 "10k":   ("Rezistory","Rezistor 10 kΩ","0,25 W, 1 %, DIN 0207","libovolný výrobce","10k","",1),
 "100k":  ("Rezistory","Rezistor 100 kΩ","0,25 W, 1 %, DIN 0207","libovolný výrobce","100k","",1),
 "100R":  ("Rezistory","Rezistor 100 Ω","0,25 W, 1 %, DIN 0207","libovolný výrobce","100R","",1),
 "0R":    ("Rezistory","Rezistor 0 Ω (propojka)","0,25 W, DIN 0207 – spojuje PGND a GND (lze nahradit drátkem)","libovolný výrobce","0R","",1),
 "BZX55C12":   ("Diody","Zenerova dioda 12 V","0,5 W, DO-35","BZX55C12 (nebo BZX79C12)","12V","pozor na polaritu – proužek = katoda (K)",2),
 "BZX55C15":   ("Diody","Zenerova dioda 15 V","0,5 W, DO-35","BZX55C15 (nebo BZX79C15)","15V","pozor na polaritu – proužek = katoda (K)",2),
 "BZX55C5V6":  ("Diody","Zenerova dioda 5,6 V","0,5 W, DO-35 – ochrana vstupů","BZX55C5V6 (nebo BZX79C5V6)","5V6","pozor na polaritu – proužek = katoda (K)",2),
 "1N4148":     ("Diody","Dioda 1N4148","DO-35","1N4148","4148","pozor na polaritu – proužek = katoda (K)",2),
 "1N4007":     ("Diody","Dioda 1N4007","1 A / 1000 V, DO-41 – flyback solenoidů","1N4007","4007","pozor na polaritu – proužek = katoda (K)",2),
 "P6KE27A":    ("Diody","Transil (TVS) 27 V","600 W, jednosměrný, DO-15","P6KE27A","P6KE27","pozor na polaritu – proužek = katoda (K)",2),
 "100nF":      ("Kondenzátory","Kondenzátor 100 nF","keramický, 50 V, X7R, rozteč 2,54 mm","např. KEMET C320C104K5R5TA","100n","",3),
 "600R@100MHz":("Kondenzátory","Feritová perla","axiální, ≥ 600 Ω @ 100 MHz, ≥ 1 A (montáž nastojato)","např. Fair-Rite 2743019447","FB","",3),
 "10uF 25V":   ("Kondenzátory","Elektrolyt 10 µF / 25 V","Ø5 mm, rozteč 2,0 mm, 105 °C","libovolný výrobce","10u","polarita! + na označený pad",6),
 "47uF 10V":   ("Kondenzátory","Elektrolyt 47 µF / 10 V (i 16 V)","Ø6,3 mm, rozteč 2,5 mm, 105 °C","libovolný výrobce","47u","polarita! + na označený pad",6),
 "100uF 10V":  ("Kondenzátory","Elektrolyt 100 µF / 10 V (i 16 V)","Ø6,3 mm, rozteč 2,5 mm, 105 °C","libovolný výrobce","100u","polarita! + na označený pad",6),
 "470uF 25V":  ("Kondenzátory","Elektrolyt 470 µF / 25 V","Ø10 mm, rozteč 5 mm, 105 °C, low-ESR","např. Panasonic řada FR","470u","polarita! + na označený pad",6),
 "PC817":      ("Polovodiče","Optočlen PC817","DIP-4","PC817C (Sharp / Everlight / Lite-On)","PC817","tečka = pin 1 (u čtvercového padu)",4),
 "BC337":      ("Polovodiče","Tranzistor NPN BC337","TO-92, pinout C-B-E (pad 1 = C)","BC337-40","BC337","plochá strana podle potisku; nožičky trochu roztáhnout (rozteč 2,54)",4),
 "zelená READY":("Polovodiče","LED zelená 3 mm","indikace READY","libovolná 3 mm","LED","polarita! kratší nožička / plochá strana = katoda (čtvercový pad)",4),
 "IRF4905":    ("Polovodiče","P-MOSFET IRF4905","TO-220 – ochrana proti přepólování","IRF4905PBF (Infineon)","IRF4905","kovová plocha podle potisku",7),
 "IRF9540N":   ("Polovodiče","P-MOSFET IRF9540N","TO-220 – spínání solenoidů","IRF9540NPBF (Infineon)","9540N","kovová plocha podle potisku",7),
 "TSR 1-2450 (12V->5V 1A)":("Polovodiče","Spínaný stabilizátor 5 V / 1 A","formát 7805 (SIP-3), vstup 6,5–36 V","Traco TSR 1-2450","TSR1-2450","pin 1 = Vin (čtvercový pad)",7),
 "G5LE-1 12V SPDT":("Mechanické","Relé 12 V SPDT","kontakty 10 A, cívka 12 V","Omron G5LE-1 DC12 (pinově kompatibilní: Songle SRD-12VDC-SL-C)","K1–K3","",8),
 "5A":         ("Mechanické","Patice pro pojistku TR5","rozteč 5,08 mm","Littelfuse 560 (TR5 patice)","F1","",7),
 "J1 +12V / PGND":("Konektory","Svorkovnice 2 pól","rozteč 5,08 mm, do DPS","Phoenix Contact MKDS 1,5/2-5,08 (1715721)","J1","otvory pro vodiče ven z desky",8),
 "J2 SOL1-3 / TCU1-3":("Konektory","Svorkovnice 6 pól","rozteč 5,08 mm, do DPS","Phoenix Contact MKDS 1,5/6-5,08 (1715763)","J2","otvory pro vodiče ven z desky",8),
 "J3 UP DOWN READY LOCKUP":("Konektory","Svorkovnice 6 pól","rozteč 5,08 mm, do DPS","Phoenix Contact MKDS 1,5/6-5,08 (1715763)","J3","otvory pro vodiče ven z desky",8),
 "J5 Nextion +5V GND RX TX":("Konektory","Konektor JST-XH 4 pin (do DPS)","rozteč 2,5 mm, vertikální","JST B4B-XH-A","J5","",5),
 "J6 rezerva":("Konektory","Kolíková lišta 1×7","rozteč 2,54 mm (volitelné – rezerva)","libovolná (odlomit z 1×40)","J6","",5),
 "JP1 5V -> Nano (jumper)":("Konektory","Kolíková lišta 1×2 + jumper","rozteč 2,54 mm","libovolná + zkratovací propojka 2,54","JP1","",5),
 "Arduino Nano":("Moduly","Arduino Nano (ATmega328P, 5 V)","zasouvá se do patic (viz níže)","Arduino Nano nebo klon s CH340","NANO","USB směrem k pravé hraně desky",9),
}
STEPS={1:"1. rezistory",2:"2. diody",3:"3. keramické C a ferity",4:"4. optočleny, tranzistor, LED",5:"5. lišty, patice, JST",6:"6. elektrolyty",7:"7. TO-220, stabilizátor, patice pojistky",8:"8. relé a svorkovnice",9:"9. Arduino Nano (zasunout)"}
EXTRA=[ # items not on the schematic, per board
 ("Konektory","Dutinková lišta 1×15 (patice pro Nano)","rozteč 2,54 mm","libovolná",2,"U2 (patice)"),
 ("Mechanické","Pojistka TR5 5 A, pomalá (T)","TR5, rozteč 5,08 mm","Littelfuse řada 372, 5 A T",1,"do F1"),
 ("Konektory","Kabelová část JST-XH 4 pin","pouzdro + 4 kontakty","JST XHP-4 + 4× SXH-001T-P0.6",1,"kabel k displeji"),
 ("Moduly","Displej Nextion 2,4″","UART, 5 V","Nextion NX3224T024 (nebo NX4024T032 3,2″)",1,"mimo desku, přes J5"),
 ("Mechanické","Distanční sloupek M3 + šroub","pro 4 díry M3 v rozích","libovolné, např. 10 mm",4,"H1–H4"),
]

groups={}
for p in parts:
    c=CAT[p["val"]]
    key=(c[0],c[1],c[2],c[3])
    groups.setdefault(key,[]).append(p["ref"])
def refkey(r):
    m=re.match(r"([A-Z]+)(\d+)",r); return (m.group(1),int(m.group(2)))
order=["Rezistory","Diody","Kondenzátory","Polovodiče","Mechanické","Konektory","Moduly"]
rows=[]
for (g,n,spec,ex),refs in groups.items():
    rows.append((g,n,spec,ex,len(refs)," ".join(sorted(refs,key=refkey))))
for g,n,spec,ex,q,refs in EXTRA: rows.append((g,n,spec,ex,q,refs))
def valsort(n):
    m=re.search(r"(\d+[,.]?\d*)\s*(k|M)?Ω",n)
    if not m: return 0
    v=float(m.group(1).replace(",","."))*{"k":1e3,"M":1e6,None:1}[m.group(2)]; return v
rows.sort(key=lambda r:(order.index(r[0]),valsort(r[1]),r[1]))

F="Arial"
thin=Side(style="thin",color="BBBBBB"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
hdr_fill=PatternFill("solid",fgColor="1F3864"); grp_fill=PatternFill("solid",fgColor="D9E1F2"); inp_fill=PatternFill("solid",fgColor="FFFF00")
wb=Workbook(); ws=wb.active; ws.title="Objednávka"
ws["A1"]="AW4 shield v0.3 (THT, 100×100 mm) – seznam součástek k objednání"; ws["A1"].font=Font(name=F,bold=True,size=14)
ws["A2"]="Počet desek:"; ws["A2"].font=Font(name=F,bold=True)
ws["C2"]=1; ws["C2"].font=Font(name=F,bold=True,color="0000FF"); ws["C2"].fill=inp_fill; ws["C2"].border=B
ws["C2"].comment=Comment("Zadejte, kolik desek budete osazovat. Sloupec 'Ks celkem' se přepočítá.","AW4")
ws["D2"]="Rezerva navíc (ks na položku, pro drobné díly):"; ws["D2"].font=Font(name=F)
ws["G2"]=2; ws["G2"].font=Font(name=F,bold=True,color="0000FF"); ws["G2"].fill=inp_fill; ws["G2"].border=B
ws["G2"].comment=Comment("Přičte se jen k rezistorům, diodám a kondenzátorům (snadno se ztratí nebo zničí).","AW4")
ws["A3"]="Žluté buňky = vstupy (modré číslo). Ostatní se počítají. Příklady dílů jsou orientační – jde o běžné typy, které se dají koupit v GM, TME, Mouser apod.; LCSC/TME kódy nejsou ověřené."
ws["A3"].font=Font(name=F,italic=True,size=9,color="555555")
H=["Pol.","Skupina","Součástka","Specifikace","Příklad dílu (výrobce / typ)","Ks / deska","Ks celkem","Označení na desce","Koupeno ✓"]
r0=5
for i,h in enumerate(H,1):
    c=ws.cell(row=r0,column=i,value=h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=hdr_fill; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
r=r0+1; pol=1; first=r
for g,n,spec,ex,q,refs in rows:
    vals=[pol,g,n,spec,ex,q,None,refs,""]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F,size=10); c.border=B; c.alignment=Alignment(vertical="top",wrap_text=True)
    small = g in ("Rezistory","Diody","Kondenzátory")
    ws.cell(row=r,column=7,value=f"=F{r}*$C$2+IF({str(small).upper()},$G$2,0)" if small else f"=F{r}*$C$2")
    ws.cell(row=r,column=7).font=Font(name=F,size=10,bold=True)
    ws.cell(row=r,column=6).alignment=Alignment(horizontal="center",vertical="top"); ws.cell(row=r,column=7).alignment=Alignment(horizontal="center",vertical="top")
    if r%2==0:
        for i in range(1,10): ws.cell(row=r,column=i).fill=PatternFill("solid",fgColor="F2F2F2")
    r+=1; pol+=1
last=r-1
ws.cell(row=r+1,column=5,value="Celkem kusů k objednání:").font=Font(name=F,bold=True)
ws.cell(row=r+1,column=7,value=f"=SUM(G{first}:G{last})").font=Font(name=F,bold=True)
ws.cell(row=r+2,column=5,value=f"Součástek osazených na desce (referencí ve schématu): {len(parts)}").font=Font(name=F,size=9,color="555555")
for i,w in enumerate([5,13,30,42,40,8,9,34,10],1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="A6"; ws.auto_filter.ref=f"A{r0}:I{last}"
ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True

# --- assembly sheet ---
ws2=wb.create_sheet("Osazení")
ws2["A1"]="Postup osazení – od nejnižších součástek po nejvyšší"; ws2["A1"].font=Font(name=F,bold=True,size=14)
ws2["A2"]="Na desce je na potisku vytištěné, co kam patří (sloupec 'Na potisku'). Označení (R18 …) je v osazovacím výkresu docs/aw4_shield_osazovaci_vykres.pdf a v interaktivním BOM."
ws2["A2"].font=Font(name=F,italic=True,size=9,color="555555")
H2=["Krok","Označení","Na potisku","Součástka","Hodnota ve schématu","Pozor","Osazeno ✓"]
for i,h in enumerate(H2,1):
    c=ws2.cell(row=4,column=i,value=h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=hdr_fill; c.border=B; c.alignment=Alignment(horizontal="center",wrap_text=True)
items=[]
for p in parts:
    g,n,spec,ex,lab,pol,step=CAT[p["val"]]
    if p["ref"][0] in "JK" or p["ref"]=="JP1": lab=p["ref"]
    items.append((step,refkey(p["ref"]),p["ref"],lab,n,p["val"],pol))
items.sort()
r=5; prev=None
for step,_,ref,lab,n,val,pol in items:
    if step!=prev:
        c=ws2.cell(row=r,column=1,value=STEPS[step]); c.font=Font(name=F,bold=True); 
        for i in range(1,8): ws2.cell(row=r,column=i).fill=grp_fill; ws2.cell(row=r,column=i).border=B
        r+=1; prev=step
    for i,v in enumerate(["",ref,lab,n,val,pol,""],1):
        c=ws2.cell(row=r,column=i,value=v); c.font=Font(name=F,size=10,color=("C00000" if (i==6 and v) else "000000")); c.border=B; c.alignment=Alignment(vertical="top",wrap_text=True)
    r+=1
for i,w in enumerate([30,10,12,32,26,48,10],1): ws2.column_dimensions[get_column_letter(i)].width=w
ws2.freeze_panes="A5"
ws2.page_setup.orientation="landscape"; ws2.sheet_properties.pageSetUpPr.fitToPage=True; ws2.page_setup.fitToWidth=1; ws2.page_setup.fitToHeight=0
wb.save(OUT); print("rows",len(rows),"parts",len(parts))
